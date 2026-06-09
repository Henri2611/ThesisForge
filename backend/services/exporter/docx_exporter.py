import re
import os
import uuid
import logging
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt, Inches, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from services.exporter.template_config import TemplateConfig, BUILTIN_TEMPLATES

logger = logging.getLogger(__name__)

EXPORT_DIR = Path(__file__).parent.parent.parent / "exports"


def _hex_to_rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


class DocxExporter:
    def __init__(self):
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    def export(self, title: str, chapters: list[dict], template_id: str | None = None, config: TemplateConfig | None = None) -> str:
        if config is None:
            config = BUILTIN_TEMPLATES.get(template_id) if template_id else None
        if config is None:
            config = TemplateConfig()

        doc = DocxDocument()

        self._apply_page_setup(doc, config)
        self._set_styles(doc, config)

        if config.cover_page:
            self._add_title_page(doc, title, config)
        else:
            self._add_title_heading(doc, title, config)

        for chapter in chapters:
            self._add_chapter(doc, chapter, config)

        filename = f"{uuid.uuid4().hex}.docx"
        filepath = EXPORT_DIR / filename
        doc.save(str(filepath))
        logger.info(f"Exported document to {filepath}")
        return str(filepath)

    def _apply_page_setup(self, doc: DocxDocument, config: TemplateConfig):
        section = doc.sections[0]
        p = config.page
        section.top_margin = Cm(p.margin_top / 10)
        section.bottom_margin = Cm(p.margin_bottom / 10)
        section.left_margin = Cm(p.margin_left / 10)
        section.right_margin = Cm(p.margin_right / 10)
        if p.orientation == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width

        if config.header:
            header = section.header
            p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            text = config.header.replace("{page_number}", "<page_number>")
            p.clear()
            run = p.add_run(text)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        if config.footer_page_numbers:
            footer = section.footer
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("— Page ")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            self._add_page_field(p)
            run2 = p.add_run(" —")
            run2.font.size = Pt(9)
            run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    def _add_page_field(self, paragraph):
        from lxml import etree
        run1 = paragraph.add_run()
        etree.SubElement(run1._r, qn("w:fldChar")).set(qn("w:fldCharType"), "begin")
        run2 = paragraph.add_run()
        instr = etree.SubElement(run2._r, qn("w:instrText"))
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        run3 = paragraph.add_run()
        etree.SubElement(run3._r, qn("w:fldChar")).set(qn("w:fldCharType"), "end")

    def _set_styles(self, doc: DocxDocument, config: TemplateConfig):
        style = doc.styles["Normal"]
        style.font.name = config.fonts.body
        style.font.size = Pt(config.fonts.size_body)
        style.paragraph_format.line_spacing = config.spacing.line
        style.paragraph_format.space_after = Pt(config.spacing.after_para)

    def _add_title_page(self, doc: DocxDocument, title: str, config: TemplateConfig):
        for _ in range(8):
            doc.add_paragraph()

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(config.fonts.size_h1 + 6)
        run.font.color.rgb = _hex_to_rgb(config.color_primary)
        run.font.name = config.fonts.heading

        doc.add_page_break()

    def _add_title_heading(self, doc: DocxDocument, title: str, config: TemplateConfig):
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(config.fonts.size_h1)
        run.font.color.rgb = _hex_to_rgb(config.color_primary)
        run.font.name = config.fonts.heading
        p.paragraph_format.space_after = Pt(config.spacing.after_para)

    def _add_chapter(self, doc: DocxDocument, chapter: dict, config: TemplateConfig):
        order = chapter.get("order", 1)
        title = chapter.get("title", f"Chapter {order}")
        content = chapter.get("content", "")

        if config.heading_numbering:
            heading_text = f"{order}. {title}"
        else:
            heading_text = title

        p = doc.add_paragraph()
        run = p.add_run(heading_text)
        run.bold = True
        run.font.size = Pt(config.fonts.size_h1)
        run.font.color.rgb = _hex_to_rgb(config.color_primary)
        run.font.name = config.fonts.heading
        p.paragraph_format.space_before = Pt(config.spacing.before_heading)
        p.paragraph_format.space_after = Pt(config.spacing.after_para)

        self._add_markdown_content(doc, content, config)

    def _add_markdown_content(self, doc: DocxDocument, content: str, config: TemplateConfig):
        lines = content.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i]

            if line.strip().startswith("```"):
                i += 1
                code_lines = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                if code_lines:
                    self._add_code_block(doc, code_lines, config)
                i += 1
                continue

            heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
            if heading_match:
                level = len(heading_match.group(1))
                text = heading_match.group(2)
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True
                run.font.name = config.fonts.heading
                size_map = {1: config.fonts.size_h1, 2: config.fonts.size_h2, 3: config.fonts.size_h3}
                run.font.size = Pt(size_map.get(level, config.fonts.size_body))
                run.font.color.rgb = _hex_to_rgb(config.color_primary)
                p.paragraph_format.space_before = Pt(config.spacing.before_heading)
                p.paragraph_format.space_after = Pt(config.spacing.after_para)
                i += 1
                continue

            if line.strip() in ("---", "***", "___"):
                doc.add_paragraph("_" * 60)
                i += 1
                continue

            list_match = re.match(r"^(\s*)[-*+]\s+(.+)$", line)
            if list_match:
                indent = len(list_match.group(1))
                text = list_match.group(2)
                p = doc.add_paragraph(style="List Bullet")
                p.clear()
                run = p.add_run(text)
                run.font.size = Pt(config.fonts.size_body)
                p.paragraph_format.left_indent = Cm(1.27 + indent * 0.63)
                i += 1
                continue

            if not line.strip():
                doc.add_paragraph()
                i += 1
                continue

            p = doc.add_paragraph()
            self._add_inline_formatting(p, line, config)
            i += 1

    def _add_inline_formatting(self, paragraph, text: str, config: TemplateConfig):
        parts = re.split(r"(\*\*.*?\*\*|__.*?__|\*.*?\*|`.*?`)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("__") and part.endswith("__"):
                run = paragraph.add_run(part[2:-2])
                run.italic = True
            elif part.startswith("*") and part.endswith("*"):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif part.startswith("`") and part.endswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = config.fonts.code
                run.font.size = Pt(config.fonts.size_body - 2)
            else:
                paragraph.add_run(part)

    def _add_code_block(self, doc: DocxDocument, code_lines: list[str], config: TemplateConfig):
        for code_line in code_lines:
            p = doc.add_paragraph()
            run = p.add_run(code_line)
            run.font.name = config.fonts.code
            run.font.size = Pt(config.fonts.size_body - 3)
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
